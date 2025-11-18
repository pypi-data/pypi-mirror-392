from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any, Dict, List, Optional

from langchain_core.documents import Document

from ..base_loader import BaseMarkitdownLoader

if TYPE_CHECKING:
    from markitdown import MarkItDown


class XlsxLoader(BaseMarkitdownLoader):
    """Loader for XLSX files with optional sheet-level splitting."""

    def __init__(
        self,
        file_path: str,
        split_by_page: bool = False,
        *,
        verbose: bool = False,
        converter: Optional["MarkItDown"] = None,
    ):
        super().__init__(file_path, verbose=verbose, converter=converter)
        self.split_by_page = split_by_page

    def load(self) -> List[Document]:
        """Load and convert XLSX files, optionally splitting each worksheet."""
        metadata: Dict[str, Any] = {
            "source": self.file_path,
            "success": False,
            "conversion_success": False,
        }

        try:
            metadata.update(self._file_metadata())
            result = self._convert_to_markdown()
            metadata["success"] = True
            metadata["conversion_success"] = True
            metadata.update(self._extract_conversion_metadata(result))
            metadata.update(self._extract_workbook_metadata())

            markdown_content = self._get_text_content(result)

            if not self.split_by_page:
                metadata["content_type"] = "workbook"
                return [Document(page_content=markdown_content, metadata=metadata)]

            documents = self._split_workbook(result, markdown_content, metadata)
            return documents
        except FileNotFoundError as exc:
            metadata["error"] = "File not found."
            raise ValueError(
                f"Markitdown conversion failed for {self.file_path}: File not found",
            ) from exc
        except Exception as exc:
            metadata["error"] = str(exc)
            raise ValueError(
                f"Failed to load and convert XLSX file: {exc}",
            ) from exc

    def _extract_workbook_metadata(self) -> Dict[str, Any]:
        """Fetch workbook metadata using openpyxl when available."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"metadata_extraction_warning": "openpyxl not installed"}

        try:
            workbook = load_workbook(self.file_path, read_only=True, data_only=True)
            props = workbook.properties
        except Exception as exc:  # pragma: no cover - relies on openpyxl internals
            return {"metadata_extraction_error": str(exc)}

        metadata: Dict[str, Any] = {}
        attr_map = {
            "creator": "author",
            "title": "title",
            "subject": "subject",
            "description": "description",
            "keywords": "keywords",
            "category": "category",
        }
        for attr, key in attr_map.items():
            value = getattr(props, attr, None)
            if value:
                metadata[key] = value

        return metadata

    def _split_workbook(
        self,
        result: Any,
        markdown_content: str,
        metadata: Dict[str, Any],
    ) -> List[Document]:
        """Split workbook content either via MarkItDown pages or a fallback parser."""
        pages = getattr(result, "pages", None)
        documents: List[Document] = []

        if isinstance(pages, list) and pages:
            for index, page_content in enumerate(pages, start=1):
                page_metadata = metadata.copy()
                page_metadata["page_number"] = index
                page_metadata["content_type"] = "worksheet"
                documents.append(
                    Document(page_content=page_content, metadata=page_metadata),
                )
            return documents

        # Fallback: attempt to split by sheet headers in the markdown output.
        pattern = re.compile(r"^##\s+(.+)$", re.MULTILINE)
        splits = pattern.split(markdown_content)

        if len(splits) <= 1:
            fallback_metadata = metadata.copy()
            fallback_metadata["content_type"] = "worksheet"
            return [Document(page_content=markdown_content, metadata=fallback_metadata)]

        # regex split returns alternating content/header pairs; rebuild them.
        leading_content = splits[0].strip()
        page_counter = 1
        if leading_content:
            documents.append(
                Document(
                    page_content=leading_content,
                    metadata={
                        **metadata,
                        "content_type": "worksheet",
                        "page_number": page_counter,
                    },
                ),
            )
            page_counter += 1

        for idx in range(1, len(splits), 2):
            sheet_name = splits[idx].strip()
            sheet_body = splits[idx + 1].strip()

            page_metadata = metadata.copy()
            page_metadata["page_number"] = page_counter
            page_metadata["worksheet"] = sheet_name
            page_metadata["content_type"] = "worksheet"

            documents.append(
                Document(page_content=sheet_body, metadata=page_metadata),
            )
            page_counter += 1

        return documents
