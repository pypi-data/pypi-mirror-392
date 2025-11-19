
import pytest
from converters.docx2csv import Docx2CsvConverter
import os
from docx import Document

# Helper function to create a dummy DOCX with a table for testing
def create_dummy_docx_with_table(file_path):
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = 'Header 1'
    table.cell(0, 1).text = 'Header 2'
    table.cell(0, 2).text = 'Header 3'
    table.cell(1, 0).text = 'Row 1 Col 1'
    table.cell(1, 1).text = 'Row 1 Col 2'
    table.cell(1, 2).text = 'Row 1 Col 3'
    table.cell(2, 0).text = 'Row 2 Col 1'
    table.cell(2, 1).text = 'Row 2 Col 2'
    table.cell(2, 2).text = 'Row 2 Col 3'
    document.save(file_path)

def test_docx2csv_conversion(tmp_path):
    """
    Tests the DOCX to CSV conversion.
    """
    input_docx = tmp_path / "dummy_table.docx"
    output_csv = tmp_path / "output.csv"
    
    create_dummy_docx_with_table(input_docx)

    converter = Docx2CsvConverter()
    converter.convert(str(input_docx), str(output_csv))

    assert os.path.exists(output_csv)
    assert os.path.getsize(output_csv) > 0

# TODO: Add tests for AST generation once implemented.
